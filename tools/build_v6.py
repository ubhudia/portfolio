#!/usr/bin/env python3
"""
build_v6.py

Builds Portfolio_Case_Study_Review_v6.xlsx from the live index.html and a plain
text export of MasterCV.

  python3 tools/build_v6.py --html index.html --mastercv mastercv.txt \
      --out tools/Portfolio_Case_Study_Review_v6.xlsx

What it does
  1. Reads SKILLS and CASES out of index.html (needs node on the PATH).
  2. Gives each card a stable UID (company slug + per-company counter, in array
     order), matching the v5 scheme (kseye-05, weareaspir-03, hays-01).
  3. Checks every figure on every card against the MasterCV section for that
     employer, then against the whole of MasterCV, and records the result.
  4. Lists MasterCV achievement lines whose figures appear nowhere on the site
     as candidate new cards. Lines marked as estimates, and lines carrying
     modelled pipeline currency figures, are excluded.
  5. Writes the workbook. It does NOT change index.html. Umesh reviews the
     workbook; a separate apply step edits CASES and regenerates the cards.

MasterCV text format: one block per line, headings prefixed with # marks, as
produced by the export step used in the October 2026 session. Section headings
look like "### 4.5 Hays, Global Head of Growth, 07/2020 to 03/2023".
"""
import argparse
import collections
import json
import re
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ---- employer name on the site -> MasterCV section number -------------------
SECTION = {
    "Eagle Eye": "4.1", "KSeye": "4.2", "Randstad": "4.4", "Hays": "4.5",
    "Marks & Spencer": "4.6", "Pomelo Pay": "4.7", "We Are Aspire": "4.8",
    "Hubbul": "4.9", "EVEN Card": "4.10", "Cigna Corporation": "4.11",
    "Family Brands": "4.12", "Ladbrokes": "4.13", "Accessories Direct": "4.14",
    "Alpari (UK)": "4.15", "Signet Group": "4.16", "LOVEFiLM UK": "4.17",
    "Euromoney plc": "4.18",
}

# Pipeline currency figures MasterCV marks as derived or modelled. They never go
# on the public site (README, "Master CV as ground truth").
MODELLED = [
    r"£6\.0m", r"£2\.5m", r"£1\.6m", r"~£352k", r"~£544k", r"~£704k",
    r"~£1\.5m in pipeline", r"\bestimate\b", r"\bmodelled\b",
]

FIGURE = re.compile(
    r"(?:£|\$|€)\d[\d,.]*\s?(?:k|m|bn|M|K)?\b"       # money
    r"|\d[\d,.]*\s?%"                                 # percentages
    r"|\b\d[\d,.]*\s?(?:x|X)\b"                       # multiples
    r"|\b\d[\d,.]*\s?(?:countries|markets|languages|brands|sites|events|"
    r"people|staff|agencies|webinars|leads|MQLs|SQLs|accounts|assets|"
    r"months|weeks|days|units|hours)\b"               # counted nouns
)

TAG_ORDER = []  # filled from SKILLS


def norm(s):
    s = s.lower().replace("–", "-").replace("—", "-")
    s = s.replace(",", "").replace(" ", "")
    return s


def load_site(html_path):
    src = open(html_path, encoding="utf-8").read().split("\n")
    i = next(n for n, l in enumerate(src) if "const SKILLS = [" in l)
    k = next(n for n, l in enumerate(src) if "const LOGO_MAP" in l)
    code = "\n".join(src[i:k]) + "\n;JSON.stringify({SKILLS,CASES})"
    out = subprocess.run(["node", "-e", f"process.stdout.write(eval({json.dumps(code)}))"],
                         capture_output=True, text=True, check=True).stdout
    return json.loads(out)


def load_mastercv(path):
    lines = open(path, encoding="utf-8").read().split("\n")
    sections, cur = collections.OrderedDict(), None
    for ln in lines:
        m = re.match(r"#+\s*(\d+\.\d+)\s", ln)
        if m:
            cur = m.group(1)
            sections.setdefault(cur, [])
            continue
        if re.match(r"#\s*§\d+\.", ln):
            cur = None
            continue
        if cur:
            sections[cur].append(ln)
    return sections, "\n".join(lines)


def figures(text):
    seen, out = set(), []
    for m in FIGURE.finditer(text):
        f = m.group(0).strip()
        if f not in seen:
            seen.add(f)
            out.append(f)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--html", default="index.html")
    ap.add_argument("--mastercv", required=True)
    ap.add_argument("--out", default="tools/Portfolio_Case_Study_Review_v6.xlsx")
    a = ap.parse_args()

    site = load_site(a.html)
    skills, cases = site["SKILLS"], site["CASES"]
    TAG_ORDER[:] = [s["key"] for s in skills]
    label = {s["key"]: s["label"] for s in skills}
    sections, whole = load_mastercv(a.mastercv)
    whole_n = norm(whole)
    site_text_n = norm(" ".join(c["headline"] + " " + c["narrative"] for c in cases))

    # ---- UIDs ----------------------------------------------------------------
    counter = collections.Counter()
    for pos, c in enumerate(cases, 1):
        slug = re.sub(r"[^a-z0-9]", "", c["company"].lower())[:10]
        counter[c["company"]] += 1
        c["uid"] = f"{slug}-{counter[c['company']]:02d}"
        c["id"] = pos

    # ---- verification --------------------------------------------------------
    for c in cases:
        sec_key = SECTION.get(c["company"])
        sec_n = norm("\n".join(sections.get(sec_key, [])))
        figs = figures(c["headline"] + " " + c["narrative"])
        in_sec, elsewhere, missing = [], [], []
        for f in figs:
            fn = norm(f)
            if fn in sec_n:
                in_sec.append(f)
            elif fn in whole_n:
                elsewhere.append(f)
            else:
                missing.append(f)
        if not figs:
            status = "No figures (wording check only)"
        elif not missing and not elsewhere:
            status = "Verified"
        elif not missing:
            status = "Check: figure found outside this employer's section"
        else:
            status = "Not found in MasterCV"
        flag = [p for p in MODELLED if re.search(p, c["headline"] + " " + c["narrative"], re.I)]
        c["status"], c["figs"] = status, figs
        c["missing"], c["elsewhere"] = missing, elsewhere
        c["section"] = f"§{sec_key}" if sec_key else ""
        c["modelled_flag"] = ", ".join(flag)

    # ---- candidate new cards ---------------------------------------------------
    cands = []
    for company, sec_key in SECTION.items():
        for ln in sections.get(sec_key, []):
            ln = ln.strip()
            if len(ln) < 60:
                continue
            if any(re.search(p, ln, re.I) for p in MODELLED):
                continue
            figs = figures(ln)
            if not figs:
                continue
            if any(norm(f) in site_text_n for f in figs):
                continue
            first = re.split(r"(?<=[a-z0-9\)])\.\s+(?=[A-Z])", ln)[0].rstrip(".") + "."
            cands.append({
                "company": company, "section": f"§{sec_key}", "figs": figs,
                "line": first if len(first) <= 300 else first[:297].rsplit(" ", 1)[0] + "...",
                "full_len": len(ln),
            })
    cands.sort(key=lambda r: (-len(r["figs"]), r["company"]))

    # ---- workbook --------------------------------------------------------------
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F2A3A")
    head_font = Font(bold=True, color="FFFFFF")
    wrap = Alignment(wrap_text=True, vertical="top")
    red = PatternFill("solid", fgColor="F8D7DA")
    amber = PatternFill("solid", fgColor="FFF3CD")
    green = PatternFill("solid", fgColor="D4EDDA")

    def sheet(title, headers, rows, widths):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for cell in ws[1]:
            cell.fill, cell.font, cell.alignment = head_fill, head_font, wrap
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = wrap
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "C2"
        return ws

    wb.remove(wb.active)

    ws = wb.create_sheet("Read Me First")
    notes = [
        "Portfolio Case Study Review v6",
        "",
        "Source: live index.html (64 cards, 16 tags) checked against MasterCV rev 408 (16 Oct 2026).",
        "This workbook proposes changes. Nothing has been applied to index.html.",
        "",
        "How to review",
        "1. Case Studies tab: filter 'Verified vs MasterCV'. Fix or delete anything not 'Verified'.",
        "2. Edit Headline, Narrative or the tag X marks directly. Put Y in Delete? to remove a card.",
        "3. Proposed New Cards tab: put Y in 'Add?' for lines you want as cards, and edit the wording.",
        "4. Return the file. UID is the matching key; ID is positional only.",
        "",
        "Rules applied (from MasterCV and the README)",
        "- No modelled pipeline currency (Hays 6.0m, Randstad 2.5m, KSeye 1.6m and derived splits).",
        "- Lines marked estimate in MasterCV are excluded from proposals.",
        "- No cross-role roll-ups, rankings or 'x of y roles' wording.",
        "- A figure is 'Verified' only if it appears in the same employer's section of MasterCV.",
        "- The check is on figures, not meaning. A verified figure can still be attached to the wrong claim.",
        "  Read each narrative against its MasterCV section before approving.",
        "- 'Not found' can be a wording mismatch rather than a wrong figure (for example '12 to 4 weeks' against",
        "  '12 weeks → 4', or 'doubled' against '2x'). Confirm against the section named in the MasterCV column.",
        "- Proposed New Cards are the first sentence of each MasterCV line, so many need rewording into a card headline.",
    ]
    for n in notes:
        ws.append([n])
    ws.column_dimensions["A"].width = 120
    ws["A1"].font = Font(bold=True, size=14)

    counts = collections.Counter(c["status"] for c in cases)
    tagc = collections.Counter(t for c in cases for t in c["tags"])
    cocnt = collections.Counter(c["company"] for c in cases)
    dash = wb.create_sheet("Dashboard")
    dash.append(["Verification status", "Cards"])
    for k, v in counts.most_common():
        dash.append([k, v])
    dash.append([])
    dash.append(["Tag", "Cards"])
    for key in TAG_ORDER:
        dash.append([label[key], tagc[key]])
        if tagc[key] <= 3:
            dash.cell(dash.max_row, 2).fill = red
    dash.append([])
    dash.append(["Employer", "Cards"])
    for k, v in cocnt.most_common():
        dash.append([k, v])
    dash.column_dimensions["A"].width = 52
    dash.column_dimensions["B"].width = 10

    headers = (["UID", "ID", "Company", "Role", "Duration", "Headline", "Narrative"]
               + [label[k] for k in TAG_ORDER]
               + ["MasterCV section", "Verified vs MasterCV", "Figures on card",
                  "Not found in MasterCV", "Found outside section", "Modelled flag",
                  "Logo", "Delete?"])
    rows = []
    for c in cases:
        rows.append([c["uid"], c["id"], c["company"], c["role"], c["duration"],
                     c["headline"], c["narrative"]]
                    + ["X" if k in c["tags"] else "" for k in TAG_ORDER]
                    + [c["section"], c["status"], ", ".join(c["figs"]),
                       ", ".join(c["missing"]), ", ".join(c["elsewhere"]),
                       c["modelled_flag"], c.get("logo", "mono badge"), ""])
    ws = sheet("Case Studies", headers, rows,
               [11, 5, 18, 22, 10, 30, 60] + [5] * len(TAG_ORDER) + [10, 26, 26, 22, 22, 16, 14, 8])
    st_col = 7 + len(TAG_ORDER) + 2
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, st_col).value
        ws.cell(r, st_col).fill = green if v == "Verified" else (amber if v and v.startswith(("Check", "No fig")) else red)

    sheet("Proposed New Cards",
          ["Add?", "Company", "MasterCV section", "Figures", "Proposed narrative (first sentence of the MasterCV line)",
           "Full line length", "Suggested tags (edit)"],
          [["", r["company"], r["section"], ", ".join(r["figs"]), r["line"], r["full_len"], ""] for r in cands],
          [6, 18, 10, 30, 90, 10, 24])

    sheet("Tag Reference", ["Key", "Label", "Description", "Cards"],
          [[s["key"], s["label"], s.get("desc", ""), tagc[s["key"]]] for s in skills],
          [14, 36, 90, 8])

    wb.save(a.out)
    print(f"cards={len(cases)} candidates={len(cands)} status={dict(counts)}")


if __name__ == "__main__":
    sys.exit(main())
